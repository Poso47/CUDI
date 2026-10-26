"""Abnahmetests A–K aus dem Master-Prompt (Abschnitt 20)."""

import sqlite3

import pytest

from conftest import FakeProvider, make_pdf, tool_call
from cudi.core.results import ANSWER, CONFIRM, ERROR, QUESTION


# --------------------------------------------------------------------------- A
def test_a_simple_chat_without_model(make_app):
    app = make_app()
    result = app.handle("Hallo CUDI.")
    assert result.kind == ANSWER
    assert result.text.startswith("Hallo")
    assert "kein Sprachmodell" in result.text  # ehrlich, statt Fähigkeiten vorzutäuschen
    assert not list(app.settings.workspace_dir.iterdir())  # keine Datei erzeugt


def test_a_simple_chat_with_model(make_app):
    fake = FakeProvider(["Hallo! Schön, dich zu sehen. Woran arbeiten wir heute?"])
    app = make_app({"fake": fake})
    result = app.handle("Hallo CUDI.")
    assert result.kind == QUESTION
    assert "Woran arbeiten wir" in result.text
    assert result.model == "fake/fake-1 (lokal)"
    # Persönlichkeit + Werkzeuge werden mitgegeben
    assert "J.A.R.V.I.S." in fake.calls[0]["system"]
    assert any(t["name"] == "write_file" for t in fake.calls[0]["tools"])


# --------------------------------------------------------------------------- B
def test_b_programming_question_asks_for_goal_and_creates_nothing(make_app):
    app = make_app()
    result = app.handle("Kannst du programmieren?")
    assert result.kind == QUESTION
    assert "Was möchtest du bauen" in result.text
    assert not result.actions
    assert not list(app.settings.workspace_dir.iterdir())


def test_b_system_prompt_requires_action_not_capability_lists(make_app):
    fake = FakeProvider(["Klar – was möchtest du bauen, und in welcher Sprache?"])
    app = make_app({"fake": fake})
    result = app.handle("Kannst du programmieren?")
    assert result.kind == QUESTION
    system = fake.calls[0]["system"]
    assert "Kannst du programmieren?" in system and "gezielt" in system
    assert "Handeln statt nur reden" in system


# --------------------------------------------------------------------------- C
def test_c_tool_use_reports_real_status(make_app):
    fake = FakeProvider([
        tool_call("write_file", "Ich lege das Skript an.", path="hallo.py", content="print('Hallo')\n"),
        "Fertig: hallo.py liegt im Arbeitsordner.",
    ])
    app = make_app({"fake": fake})
    first = app.handle("Schreib mir ein Hallo-Welt-Skript in Python und speichere es.")
    assert first.kind == CONFIRM  # Stufe 4 braucht Freigabe
    assert not (app.settings.workspace_dir / "hallo.py").exists()

    done = app.resolve(first.pending_id, approved=True)
    assert (app.settings.workspace_dir / "hallo.py").read_text() == "print('Hallo')\n"
    assert done.kind == ANSWER
    assert [a.status for a in done.actions] == ["ausgeführt"]
    # Das Modell hat das echte Werkzeugergebnis bekommen
    tool_msg = fake.calls[1]["messages"][-1]
    assert tool_msg["role"] == "tool" and "Datei geschrieben" in tool_msg["content"]


def test_c_failed_tool_is_never_reported_as_success(make_app):
    fake = FakeProvider([tool_call("read_file", path="gibt_es_nicht.txt"), "Erledigt!"])
    app = make_app({"fake": fake})
    result = app.handle("Lies gibt_es_nicht.txt")
    assert result.actions[0].status == "fehlgeschlagen"
    assert "fehlgeschlagen" in result.render()


def test_c_read_tools_run_without_confirmation(make_app):
    fake = FakeProvider([tool_call("list_files"), "Im Arbeitsordner liegt notiz.txt."])
    app = make_app({"fake": fake})
    (app.settings.workspace_dir / "notiz.txt").write_text("x")
    result = app.handle("Was liegt in meinem Arbeitsordner?")
    assert result.kind == ANSWER
    assert result.actions[0].ok and "notiz.txt" in fake.calls[1]["messages"][-1]["content"]


# --------------------------------------------------------------------------- D
def test_d_memory_survives_restart(make_app, tmp_path):
    app = make_app()
    stored = app.handle("Merke dir, dass mein Lieblingseditor VS Code ist")
    assert stored.kind == ANSWER and "Mein Lieblingseditor ist VS Code" in stored.text
    app.close()

    again = make_app(base=tmp_path)
    recall = again.handle("Was ist mein Lieblingseditor?")
    assert "VS Code" in recall.text
    assert recall.sources == ["Erinnerung #1"]


def test_d_memory_reaches_model_context(make_app):
    fake = FakeProvider(["Du nutzt VS Code."])
    app = make_app({"fake": fake})
    app.handle("Merke dir: Mein Lieblingseditor ist VS Code")
    app.handle("Welchen Editor nutze ich am liebsten?")
    prompt = fake.calls[0]["messages"][-1]["content"]
    assert "<kontext>" in prompt and "Mein Lieblingseditor ist VS Code" in prompt


# --------------------------------------------------------------------------- E
def test_e_project_state_survives_restart(make_app, tmp_path):
    app = make_app()
    app.handle("/projekt neu CUDI | Persönlicher Assistent")
    app.handle("/entscheidung SQLite als zentrale Datenbank")
    app.handle("/schritt Model Manager anbinden")
    app.close()

    fake = FakeProvider(["Wir waren beim Model Manager."])
    again = make_app({"fake": fake}, base=tmp_path)
    status = again.handle("/projektstand")
    assert "SQLite als zentrale Datenbank" in status.text and "Model Manager anbinden" in status.text
    again.handle("Wo waren wir stehen geblieben?")
    assert "Aktives Projekt" in fake.calls[0]["messages"][-1]["content"]
    assert "Model Manager anbinden" in fake.calls[0]["messages"][-1]["content"]


def test_e_conversation_history_survives_restart(make_app, tmp_path):
    fake = FakeProvider(["Notiert."])
    app = make_app({"fake": fake})
    app.handle("Wir bauen heute den Dokumentenimport.")
    app.close()
    fake2 = FakeProvider(["Weiter geht's."])
    again = make_app({"fake": fake2}, base=tmp_path)
    again.handle("Weiter?")
    contents = [m["content"] for m in fake2.calls[0]["messages"]]
    assert "Wir bauen heute den Dokumentenimport." in contents


# --------------------------------------------------------------------------- F
def test_f_document_knowledge_with_source(make_app, tmp_path):
    import docx

    app = make_app()
    doc = docx.Document()
    doc.add_heading("Lager", level=1)
    doc.add_paragraph("Der Zugangscode für das Lager lautet 4711-ALPHA.")
    docx_path = tmp_path / "handbuch.docx"
    doc.save(docx_path)
    pdf_path = make_pdf(tmp_path / "vertrag.pdf", ["Deckblatt", "Die Kuendigungsfrist betraegt 3 Monate."])

    assert app.handle(f"/import {docx_path}").kind == ANSWER
    assert "2 Abschnitte" in app.handle(f"/import {pdf_path}").text

    answer = app.handle("Wie lautet der Zugangscode für das Lager?")
    assert "4711-ALPHA" in answer.text
    assert any("handbuch.docx" in s for s in answer.sources)

    pdf_answer = app.handle("/suche Kuendigungsfrist")
    assert "3 Monate" in pdf_answer.text and "vertrag.pdf, Seite 2" in pdf_answer.sources

    duplicate = app.handle(f"/import {docx_path}")
    assert "bereits importiert" in duplicate.text
    assert len(app.knowledge.documents()) == 2


def test_f_document_excerpts_reach_model_as_data(make_app, tmp_path):
    fake = FakeProvider(["Laut handbuch.txt: 4711-ALPHA."])
    app = make_app({"fake": fake})
    (app.settings.import_dir / "handbuch.txt").write_text("Der Zugangscode für das Lager lautet 4711-ALPHA.")
    app.handle("/importordner")
    result = app.handle("Wie lautet der Zugangscode für das Lager?")
    prompt = fake.calls[0]["messages"][-1]["content"]
    assert '<dokument name="handbuch.txt"' in prompt and "4711-ALPHA" in prompt
    assert any("handbuch.txt" in s for s in result.sources)


# --------------------------------------------------------------------------- G
def test_g_missing_model_is_reported(make_app):
    app = make_app({"fake": FakeProvider(available=False)})
    result = app.handle("Erklär mir bitte Quantencomputer.")
    assert "Sprachmodell" in result.text
    report = app.handle("/status").text
    assert "kein Sprachmodell verfügbar" in report


def test_g_fallback_to_next_provider(make_app):
    broken = FakeProvider(name="erst", error="Verbindung abgelehnt")
    working = FakeProvider(["Antwort vom Ersatzmodell"], name="zweit")
    app = make_app({"erst": broken, "zweit": working}, provider_order=["erst", "zweit"])
    result = app.handle("Test")
    assert result.text == "Antwort vom Ersatzmodell"
    assert result.model.startswith("zweit/")
    calls = app.db.query("SELECT provider, ok, error FROM model_calls ORDER BY id")
    assert [(r["provider"], r["ok"]) for r in calls] == [("erst", 0), ("zweit", 1)]


def test_g_all_providers_failing_gives_clear_error(make_app):
    app = make_app({"fake": FakeProvider(error="Timeout")})
    result = app.handle("Test")
    assert result.kind == ERROR
    assert "Timeout" in result.text and "Ollama" in result.text


# --------------------------------------------------------------------------- H
def test_h_conflicting_information_triggers_question(make_app):
    app = make_app()
    app.handle("Merke dir: Mein Lieblingseditor ist VS Code")
    conflict = app.handle("Merke dir: Mein Lieblingseditor ist PyCharm")
    assert conflict.kind == QUESTION
    assert "VS Code" in conflict.text and "PyCharm" in conflict.text
    assert app.handle("ja").kind == ANSWER
    assert "PyCharm" in app.handle("Was ist mein Lieblingseditor?").text
    old = app.memory.get(1)
    assert old.status == "outdated"  # nicht gelöscht, nachvollziehbar


def test_h_conflict_rejected_keeps_old(make_app):
    app = make_app()
    app.handle("Merke dir: Ich wohne in Berlin")
    app.handle("Merke dir: Ich wohne in Hamburg")
    app.handle("nein")
    assert "Berlin" in app.handle("Wo wohne ich?").text
    assert app.memory.count() == 1


# --------------------------------------------------------------------------- I
def test_i_protected_action_needs_confirmation_and_can_be_denied(make_app):
    fake = FakeProvider([tool_call("write_file", path="x.txt", content="neu"),
                         "Verstanden, ich schreibe die Datei nicht."])
    app = make_app({"fake": fake})
    first = app.handle("Schreib x.txt")
    assert first.kind == CONFIRM and "x.txt" in first.text
    denied = app.resolve(first.pending_id, approved=False)
    assert not (app.settings.workspace_dir / "x.txt").exists()
    assert denied.actions[-1].status == "abgelehnt"


def test_i_typed_yes_confirms_and_writing_outside_workspace_is_blocked(make_app, tmp_path):
    fake = FakeProvider([tool_call("write_file", path=str(tmp_path / "ausserhalb.txt"), content="x"), "Ok."])
    app = make_app({"fake": fake})
    assert app.handle("Schreib die Datei").kind == CONFIRM
    result = app.handle("ja")
    assert not (tmp_path / "ausserhalb.txt").exists()
    assert result.actions[-1].status == "fehlgeschlagen"


def test_i_deleting_memory_requires_confirmation(make_app):
    app = make_app()
    app.handle("Merke dir: Mein Auto ist blau")
    ask = app.handle("Vergiss #1")
    assert ask.kind == CONFIRM
    assert app.memory.get(1) is not None
    app.resolve(ask.pending_id, approved=True)
    assert app.memory.get(1) is None


# --------------------------------------------------------------------------- J
def test_j_existing_modules_still_import():
    import cudi_brain
    import cudi_command_router
    import cudi_conversation_handler
    import cudi_event_bus
    import cudi_telemetry

    assert hasattr(cudi_brain, "CudiBrain")
    assert hasattr(cudi_event_bus, "CUDIEventBus")
    assert hasattr(cudi_command_router, "CUDICommandRouter")
    assert hasattr(cudi_conversation_handler, "CUDIConversationHandler")
    assert hasattr(cudi_telemetry, "get_telemetry")


def test_j_gui_keeps_old_methods_and_gains_core_integration():
    pytest.importorskip("PySide6.QtWidgets", exc_type=ImportError)
    import cudi_supreme_gui as gui

    window = gui.CUDISupremeMainWindow
    for name in ("process_user_input", "process_user_input_no_questions_original", "_create_real_files",
                 "trigger_research_mode", "open_plugin_manager", "_send_to_core"):
        assert hasattr(window, name), name


# --------------------------------------------------------------------------- K
def test_k_restart_keeps_data_consistent_without_duplicates(make_app, tmp_path):
    app = make_app()
    app.handle("Merke dir: Mein Lieblingseditor ist VS Code")
    app.handle("/aufgabe Tests schreiben")
    app.close()

    again = make_app(base=tmp_path)
    assert again.handle("Merke dir: Mein Lieblingseditor ist VS Code").text.startswith("Das weiß ich schon")
    assert again.memory.count() == 1
    assert len(again.projects.tasks()) == 1
    assert again.db.integrity_ok()


def test_k_backup_and_restore(make_app):
    app = make_app()
    app.handle("Merke dir: Mein Lieblingseditor ist VS Code")
    backup = app.db.backup(app.settings.backup_dir)
    app.handle("Merke dir: Mein Hund heißt Bello")
    assert app.memory.count() == 2

    ask = app.handle(f"/wiederherstellen {backup.name}")
    assert ask.kind == CONFIRM
    done = app.resolve(ask.pending_id, approved=True)
    assert done.kind == ANSWER
    assert app.memory.count() == 1
    with sqlite3.connect(backup) as conn:
        assert conn.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
