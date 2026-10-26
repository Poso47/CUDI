"""Wissensdatenbank: Dokumente importieren, in Abschnitte teilen, durchsuchen.

Importierte Inhalte sind Wissensquellen, keine Anweisungen. Sie werden im Prompt
als Daten markiert (siehe context_builder).
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import logging
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Callable, Dict, List, Tuple

from .database import Database, fts_query, now_iso, search_terms

log = logging.getLogger("cudi.brain.knowledge")

CHUNK_CHARS = 1200
TEXT_TYPES = {".txt", ".md", ".markdown", ".log", ".py", ".rst"}


class ImportErrorCUDI(Exception):
    """Dokument kann nicht importiert werden (Nutzer-verständliche Meldung)."""


@dataclass
class Section:
    location: str
    text: str


@dataclass
class Hit:
    document: str
    location: str
    text: str
    document_id: int

    @property
    def source(self) -> str:
        return f"{self.document}, {self.location}"


# ---------------------------------------------------------------- Extraktion
def _split_long(text: str, location: str, limit: int = CHUNK_CHARS) -> List[Section]:
    text = text.strip()
    if not text:
        return []
    if len(text) <= limit:
        return [Section(location, text)]
    parts, current = [], ""
    for para in re.split(r"\n\s*\n", text):
        para = para.strip()
        if not para:
            continue
        while len(para) > limit:  # sehr lange Absätze hart teilen
            cut = para.rfind(" ", 0, limit)
            cut = cut if cut > limit // 2 else limit
            if current:
                parts.append(current)
                current = ""
            parts.append(para[:cut])
            para = para[cut:].strip()
        if len(current) + len(para) + 2 > limit and current:
            parts.append(current)
            current = para
        else:
            current = f"{current}\n\n{para}" if current else para
    if current:
        parts.append(current)
    if len(parts) == 1:
        return [Section(location, parts[0])]
    return [Section(f"{location}, Teil {i}", p) for i, p in enumerate(parts, 1)]


def extract_pdf(path: Path) -> List[Section]:
    try:
        from pypdf import PdfReader
    except ImportError as e:
        raise ImportErrorCUDI("PDF-Import benötigt das Paket 'pypdf' (pip install pypdf).") from e
    reader = PdfReader(str(path))
    sections: List[Section] = []
    for number, page in enumerate(reader.pages, 1):
        sections += _split_long(page.extract_text() or "", f"Seite {number}")
    return sections


def extract_docx(path: Path) -> List[Section]:
    try:
        import docx  # python-docx
    except ImportError as e:
        raise ImportErrorCUDI("DOCX-Import benötigt das Paket 'python-docx' (pip install python-docx).") from e
    document = docx.Document(str(path))
    sections: List[Section] = []
    heading, buffer, counter = None, [], 0

    def flush():
        nonlocal buffer, counter
        if buffer:
            counter += 1
            label = f"Abschnitt {counter}" + (f" „{heading}“" if heading else "")
            sections.extend(_split_long("\n".join(buffer), label))
            buffer = []

    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower() if para.style is not None else ""
        if style.startswith("heading") or style.startswith("überschrift"):
            flush()
            heading = text
            buffer.append(text)
        else:
            buffer.append(text)
    flush()
    for t_index, table in enumerate(document.tables, 1):
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        sections += _split_long("\n".join(rows), f"Tabelle {t_index}")
    return sections


def extract_text(path: Path) -> List[Section]:
    text = path.read_text(encoding="utf-8", errors="replace")
    sections: List[Section] = []
    if path.suffix.lower() in (".md", ".markdown"):
        blocks = re.split(r"(?m)^(?=#{1,6}\s)", text)
        for i, block in enumerate(blocks, 1):
            title = block.strip().splitlines()[0].lstrip("#").strip() if block.strip() else ""
            sections += _split_long(block, f"Abschnitt {i}" + (f" „{title}“" if title else ""))
        return sections
    lines = text.splitlines()
    start = 0
    while start < len(lines):
        size, end = 0, start
        while end < len(lines) and size + len(lines[end]) < CHUNK_CHARS:
            size += len(lines[end]) + 1
            end += 1
        end = max(end, start + 1)
        chunk = "\n".join(lines[start:end])
        if chunk.strip():
            sections.append(Section(f"Zeilen {start + 1}–{end}", chunk))
        start = end
    return sections


def extract_json(path: Path) -> List[Section]:
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as e:
        raise ImportErrorCUDI(f"Ungültiges JSON: {e}") from e
    sections: List[Section] = []
    if isinstance(data, dict):
        for key, value in data.items():
            sections += _split_long(f"{key}: {json.dumps(value, ensure_ascii=False, indent=1)}",
                                    f"Schlüssel „{key}“")
    elif isinstance(data, list):
        for i, item in enumerate(data, 1):
            sections += _split_long(json.dumps(item, ensure_ascii=False, indent=1), f"Eintrag {i}")
    else:
        sections.append(Section("Inhalt", json.dumps(data, ensure_ascii=False)))
    return sections


def extract_csv(path: Path) -> List[Section]:
    text = path.read_text(encoding="utf-8", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:2048])
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(text), dialect))
    if not rows:
        return []
    header, body = rows[0], rows[1:]
    sections: List[Section] = []
    batch: List[str] = []
    first = 2
    for line_no, row in enumerate(body, 2):
        batch.append("; ".join(f"{h}={v}" for h, v in zip(header, row)))
        if sum(len(b) for b in batch) > CHUNK_CHARS:
            sections.append(Section(f"Zeilen {first}–{line_no}", "\n".join(batch)))
            batch, first = [], line_no + 1
    if batch:
        sections.append(Section(f"Zeilen {first}–{first + len(batch) - 1}", "\n".join(batch)))
    return sections


EXTRACTORS: Dict[str, Callable[[Path], List[Section]]] = {
    ".pdf": extract_pdf,
    ".docx": extract_docx,
    ".json": extract_json,
    ".csv": extract_csv,
}
for _suffix in TEXT_TYPES:
    EXTRACTORS[_suffix] = extract_text

SUPPORTED = sorted(EXTRACTORS)


# --------------------------------------------------------------- Speicher
class KnowledgeStore:
    def __init__(self, db: Database):
        self.db = db

    def import_file(self, path: Path) -> Tuple[str, Dict]:
        """Importiert eine Datei. Rückgabe: ('imported'|'duplicate', Dokument-Info)."""
        path = Path(path).expanduser()
        if not path.is_file():
            raise ImportErrorCUDI(f"Datei nicht gefunden: {path}")
        suffix = path.suffix.lower()
        extractor = EXTRACTORS.get(suffix)
        if not extractor:
            raise ImportErrorCUDI(
                f"Dateityp '{suffix or 'ohne Endung'}' wird nicht unterstützt. Unterstützt: {', '.join(SUPPORTED)}"
            )
        raw = path.read_bytes()
        digest = hashlib.sha256(raw).hexdigest()
        existing = self.db.query_one("SELECT * FROM documents WHERE content_hash = ?", (digest,))
        if existing:
            return "duplicate", dict(existing)

        try:
            sections = extractor(path)
        except ImportErrorCUDI:
            raise
        except Exception as e:
            raise ImportErrorCUDI(f"Inhalt von {path.name} konnte nicht gelesen werden: {e}") from e
        if not sections:
            raise ImportErrorCUDI(
                f"In {path.name} wurde kein lesbarer Text gefunden (z. B. eingescanntes PDF ohne Texterkennung)."
            )

        with self.db.transaction() as conn:
            cur = conn.execute(
                "INSERT INTO documents (path, name, file_type, content_hash, size, chunk_count, imported_at) "
                "VALUES (?, ?, ?, ?, ?, ?, ?)",
                (str(path.resolve()), path.name, suffix.lstrip("."), digest, len(raw), len(sections), now_iso()),
            )
            doc_id = cur.lastrowid
            conn.executemany(
                "INSERT INTO chunks (document_id, ordinal, location, text) VALUES (?, ?, ?, ?)",
                [(doc_id, i, s.location, s.text) for i, s in enumerate(sections)],
            )
        log.info("Dokument importiert: %s (%d Abschnitte)", path.name, len(sections))
        return "imported", dict(self.db.query_one("SELECT * FROM documents WHERE id = ?", (doc_id,)))

    def documents(self) -> List[Dict]:
        return [dict(r) for r in self.db.query("SELECT * FROM documents ORDER BY imported_at DESC")]

    def remove(self, document_id: int) -> bool:
        with self.db.transaction() as conn:
            conn.execute("DELETE FROM chunks WHERE document_id = ?", (document_id,))
            return conn.execute("DELETE FROM documents WHERE id = ?", (document_id,)).rowcount > 0

    def search(self, text: str, limit: int = 4) -> List[Hit]:
        query = fts_query(text)
        if not query:
            return []
        rows = []
        if self.db.fts_enabled:
            rows = self.db.query(
                """SELECT c.id, c.text, c.location, d.name, d.id AS doc_id
                   FROM chunks_fts f JOIN chunks c ON c.id = f.rowid JOIN documents d ON d.id = c.document_id
                   WHERE chunks_fts MATCH ? ORDER BY bm25(chunks_fts) LIMIT ?""",
                (query, limit),
            )
        if len(rows) < limit:
            terms = search_terms(text)
            seen = {r["id"] for r in rows}
            like = " OR ".join("c.text LIKE ?" for _ in terms)
            extra = self.db.query(
                f"""SELECT c.id, c.text, c.location, d.name, d.id AS doc_id
                    FROM chunks c JOIN documents d ON d.id = c.document_id WHERE {like} LIMIT ?""",
                [f"%{t}%" for t in terms] + [limit],
            )
            rows = (list(rows) + [r for r in extra if r["id"] not in seen])[:limit]
        return [Hit(r["name"], r["location"], r["text"], r["doc_id"]) for r in rows]

    @staticmethod
    def snippet(hit: Hit, query: str, width: int = 300) -> str:
        text = re.sub(r"\s+", " ", hit.text).strip()
        terms = search_terms(query)
        pos = -1
        for t in terms:
            pos = text.lower().find(t)
            if pos >= 0:
                break
        if pos < 0 or len(text) <= width:
            return text[:width] + ("…" if len(text) > width else "")
        start = max(0, pos - width // 3)
        return ("…" if start else "") + text[start:start + width] + ("…" if start + width < len(text) else "")
